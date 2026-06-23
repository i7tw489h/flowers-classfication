from __future__ import annotations

import sys
from pathlib import Path

THIS_DIR = Path(__file__).resolve().parent
VENDOR_DIR = THIS_DIR / "vendor"
if str(VENDOR_DIR) not in sys.path:
    sys.path.insert(0, str(VENDOR_DIR))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Cm, Pt


OUTPUT_DIR = THIS_DIR.parent.parent / "outputs"
OUTPUT_PPTX = OUTPUT_DIR / "deep_learning_network_architecture_cbam_editable_v2.pptx"

SLIDE_W_CM = 33.87
SLIDE_H_CM = 19.05

MODULE_H = 2.5
LAYER_W = 2.5
WIDE_W = 3.5
CBAM_W = 2.0
HEAD_W = 2.5
INPUT_W = 2.8
GAP_CM = 0.48
RADIUS_ADJ = 0.14


def rgb(value: str) -> RGBColor:
    value = value.replace("#", "")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_solid_fill(shape, color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)


def set_line(shape, color: str, width_pt: float = 1.0, dash: str | None = None) -> None:
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width_pt)
    if dash:
        shape.line.dash_style = dash


def set_text(shape, text: str, size_pt: float = 10, color: str = "#111111") -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    lines = text.split("\n")
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = line
        run.font.name = "Arial"
        run.font.bold = True
        run.font.size = Pt(size_pt)
        run.font.color.rgb = rgb(color)


def add_round_rect(slide, x: float, y: float, w: float, h: float,
                   fill: str, line: str, text: str, text_color: str,
                   text_size: float = 10, line_width: float = 1.0):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Cm(x),
        Cm(y),
        Cm(w),
        Cm(h),
    )
    set_solid_fill(shape, fill)
    set_line(shape, line, line_width)
    shape.shadow.inherit = False
    if shape.adjustments:
        shape.adjustments[0] = RADIUS_ADJ
    set_text(shape, text, text_size, text_color)
    return shape


def add_textbox(slide, x: float, y: float, w: float, h: float, text: str,
                size_pt: float, color: str = "#111111", bold: bool = True,
                align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    shape.fill.background()
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = rgb(color)
    return shape


def add_triangle_arrow_end(line_element) -> None:
    head = line_element.find(qn("a:headEnd"))
    if head is None:
        head = OxmlElement("a:headEnd")
        line_element.append(head)
    head.set("type", "triangle")
    head.set("w", "med")
    head.set("len", "med")


def add_straight_arrow(slide, x1: float, y1: float, x2: float, y2: float,
                       color: str = "#000000", width_pt: float = 1.0):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Cm(x1),
        Cm(y1),
        Cm(x2),
        Cm(y2),
    )
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(width_pt)
    conn.shadow.inherit = False
    add_triangle_arrow_end(conn._element.spPr.get_or_add_ln())
    return conn


def add_polyline_arrow(slide, points: list[tuple[float, float]], color: str = "#000000", width_pt: float = 1.0):
    connectors = []
    for i in range(len(points) - 1):
        x1, y1 = points[i]
        x2, y2 = points[i + 1]
        conn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Cm(x1),
            Cm(y1),
            Cm(x2),
            Cm(y2),
        )
        conn.line.color.rgb = rgb(color)
        conn.line.width = Pt(width_pt)
        conn.shadow.inherit = False
        if i == len(points) - 2:
            add_triangle_arrow_end(conn._element.spPr.get_or_add_ln())
        connectors.append(conn)
    return connectors


def center_y(shape) -> float:
    return shape.top.cm + shape.height.cm / 2


def right_x(shape) -> float:
    return shape.left.cm + shape.width.cm


def left_x(shape) -> float:
    return shape.left.cm


def add_operator_circle(slide, x: float, y: float, label: str,
                        fill: str = "#FFFFFF", line: str = "#666666", text_color: str = "#333333"):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL,
        Cm(x - 0.35),
        Cm(y - 0.35),
        Cm(0.7),
        Cm(0.7),
    )
    set_solid_fill(shape, fill)
    set_line(shape, line, 1.0)
    shape.shadow.inherit = False
    set_text(shape, label, 10, text_color)
    return shape


def add_group_frame(slide, x: float, y: float, w: float, h: float, fill: str, line: str):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Cm(x),
        Cm(y),
        Cm(w),
        Cm(h),
    )
    set_solid_fill(shape, fill)
    set_line(shape, line, 1.0)
    shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    shape.shadow.inherit = False
    if shape.adjustments:
        shape.adjustments[0] = 0.09
    shape.text_frame.clear()
    return shape


def build_architecture_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("#FFFFFF")

    add_textbox(slide, 0.9, 0.75, 10.0, 0.8, "Network Architecture", 18, "#111111", True, PP_ALIGN.LEFT)

    y = (SLIDE_H_CM - MODULE_H) / 2 + 0.55

    specs = [
        ("Input\n224x224x3", INPUT_W, "#FFFFFF", "#666666", "#222222"),
        ("Conv1\n7x7,s2,BN,ReLU,\nMaxPool", WIDE_W, "#F5F5F5", "#333333", "#222222"),
        ("Layer1\n256 ch", LAYER_W, "#4A90D9", "#4A90D9", "#FFFFFF"),
        ("CBAM", CBAM_W, "#F5A623", "#F5A623", "#FFFFFF"),
        ("Layer2\n512 ch", LAYER_W, "#3B7DD8", "#3B7DD8", "#FFFFFF"),
        ("CBAM", CBAM_W, "#F5A623", "#F5A623", "#FFFFFF"),
        ("Layer3\n1024 ch", LAYER_W, "#2C6AB7", "#2C6AB7", "#FFFFFF"),
        ("CBAM", CBAM_W, "#F5A623", "#F5A623", "#FFFFFF"),
        ("Layer4\n2048 ch", LAYER_W, "#1D5796", "#1D5796", "#FFFFFF"),
        ("CBAM", CBAM_W, "#F5A623", "#F5A623", "#FFFFFF"),
        ("GAP", HEAD_W, "#7ED321", "#7ED321", "#FFFFFF"),
        ("Dropout\n0.4", HEAD_W, "#7ED321", "#7ED321", "#FFFFFF"),
        ("FC", HEAD_W, "#7ED321", "#7ED321", "#FFFFFF"),
        ("102\nClasses", HEAD_W, "#7ED321", "#7ED321", "#FFFFFF"),
    ]

    total_width = sum(item[1] for item in specs) + GAP_CM * (len(specs) - 1)
    start_x = (SLIDE_W_CM - total_width) / 2

    shapes = []
    x = start_x
    for text, width, fill, line, text_color in specs:
        font_size = 10
        shape = add_round_rect(slide, x, y, width, MODULE_H, fill, line, text, text_color, font_size, 1.0)
        shapes.append(shape)
        x += width + GAP_CM

    for left_shape, right_shape in zip(shapes[:-1], shapes[1:]):
        add_straight_arrow(
            slide,
            right_x(left_shape) + 0.10,
            center_y(left_shape),
            left_x(right_shape) - 0.10,
            center_y(right_shape),
            "#000000",
            1.0,
        )


def build_cbam_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("#FFFFFF")

    add_textbox(slide, 0.9, 0.75, 8.0, 0.8, "CBAM Internal Structure", 18, "#111111", True, PP_ALIGN.LEFT)

    add_group_frame(slide, 2.0, 2.2, 13.7, 12.8, "#FFF5E6", "#E0A85B")
    add_group_frame(slide, 18.1, 2.2, 13.7, 12.8, "#F0FFF0", "#8DBB8D")

    add_textbox(slide, 5.2, 2.55, 7.2, 0.7, "Channel Attention", 14, "#8A4F08", True, PP_ALIGN.CENTER)
    add_textbox(slide, 21.1, 2.55, 7.6, 0.7, "Spatial Attention", 14, "#2D6B2D", True, PP_ALIGN.CENTER)

    input_c = add_round_rect(slide, 2.9, 7.35, 2.9, 1.5, "#FFFFFF", "#666666", "Input", "#222222", 10, 1.0)
    avg_c = add_round_rect(slide, 6.6, 5.1, 2.5, 1.5, "#FFFFFF", "#C47F1A", "AvgPool", "#444444", 10, 1.0)
    max_c = add_round_rect(slide, 6.6, 9.45, 2.5, 1.5, "#FFFFFF", "#C47F1A", "MaxPool", "#444444", 10, 1.0)
    mlp_c = add_round_rect(slide, 10.5, 7.35, 2.5, 1.5, "#FBE3BF", "#C47F1A", "Shared MLP", "#7A4A0B", 10, 1.0)
    add_node = add_operator_circle(slide, 14.3, 8.1, "+", "#FFFFFF", "#C47F1A", "#7A4A0B")
    sigmoid_c = add_round_rect(slide, 15.2, 7.35, 2.5, 1.5, "#FBE3BF", "#C47F1A", "Sigmoid", "#7A4A0B", 10, 1.0)
    mul_c = add_operator_circle(slide, 18.5, 8.1, "x", "#FFFFFF", "#C47F1A", "#7A4A0B")
    out_c = add_round_rect(slide, 19.7, 7.35, 2.9, 1.5, "#FFFFFF", "#666666", "Refined", "#222222", 10, 1.0)

    add_polyline_arrow(slide, [(5.8, 8.1), (6.2, 8.1), (6.2, 5.85), (6.6, 5.85)], "#000000", 1.0)
    add_polyline_arrow(slide, [(5.8, 8.1), (6.2, 8.1), (6.2, 10.2), (6.6, 10.2)], "#000000", 1.0)
    add_polyline_arrow(slide, [(9.1, 5.85), (9.8, 5.85), (9.8, 7.85), (10.5, 7.85)], "#000000", 1.0)
    add_polyline_arrow(slide, [(9.1, 10.2), (9.8, 10.2), (9.8, 8.35), (10.5, 8.35)], "#000000", 1.0)
    add_straight_arrow(slide, 13.0, 8.1, 13.95, 8.1, "#000000", 1.0)
    add_straight_arrow(slide, 14.65, 8.1, 15.2, 8.1, "#000000", 1.0)
    add_straight_arrow(slide, 17.7, 8.1, 18.15, 8.1, "#000000", 1.0)
    add_straight_arrow(slide, 18.85, 8.1, 19.7, 8.1, "#000000", 1.0)

    input_s = add_round_rect(slide, 18.9, 7.35, 2.9, 1.5, "#FFFFFF", "#666666", "Input", "#222222", 10, 1.0)
    avg_s = add_round_rect(slide, 22.6, 5.1, 2.5, 1.5, "#FFFFFF", "#6FA96F", "AvgPool", "#355F35", 10, 1.0)
    max_s = add_round_rect(slide, 22.6, 9.45, 2.5, 1.5, "#FFFFFF", "#6FA96F", "MaxPool", "#355F35", 10, 1.0)
    concat_s = add_round_rect(slide, 26.5, 7.35, 2.5, 1.5, "#DFF4DF", "#6FA96F", "Concat", "#355F35", 10, 1.0)
    conv_s = add_round_rect(slide, 29.8, 7.35, 2.5, 1.5, "#DFF4DF", "#6FA96F", "7x7 Conv", "#355F35", 10, 1.0)
    sigmoid_s = add_round_rect(slide, 33.1, 7.35, 2.5, 1.5, "#DFF4DF", "#6FA96F", "Sigmoid", "#355F35", 10, 1.0)
    mul_s = add_operator_circle(slide, 36.3, 8.1, "x", "#FFFFFF", "#6FA96F", "#355F35")
    out_s = add_round_rect(slide, 37.4, 7.35, 2.9, 1.5, "#FFFFFF", "#666666", "Output", "#222222", 10, 1.0)

    add_polyline_arrow(slide, [(21.8, 8.1), (22.2, 8.1), (22.2, 5.85), (22.6, 5.85)], "#000000", 1.0)
    add_polyline_arrow(slide, [(21.8, 8.1), (22.2, 8.1), (22.2, 10.2), (22.6, 10.2)], "#000000", 1.0)
    add_polyline_arrow(slide, [(25.1, 5.85), (25.8, 5.85), (25.8, 7.85), (26.5, 7.85)], "#000000", 1.0)
    add_polyline_arrow(slide, [(25.1, 10.2), (25.8, 10.2), (25.8, 8.35), (26.5, 8.35)], "#000000", 1.0)
    add_straight_arrow(slide, 29.0, 8.1, 29.8, 8.1, "#000000", 1.0)
    add_straight_arrow(slide, 32.3, 8.1, 33.1, 8.1, "#000000", 1.0)
    add_straight_arrow(slide, 35.6, 8.1, 35.95, 8.1, "#000000", 1.0)
    add_straight_arrow(slide, 36.65, 8.1, 37.4, 8.1, "#000000", 1.0)


def build_cbam_slide_refined(prs: Presentation) -> None:
    slide = prs.slides[-1]
    for shape in list(slide.shapes):
        pass


def rebuild_cbam_slide(prs: Presentation) -> None:
    prs.slides.remove(prs.slides[-1]._element)


def build_cbam_slide_final(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("#FFFFFF")

    add_textbox(slide, 0.9, 0.75, 8.0, 0.8, "CBAM Internal Structure", 18, "#111111", True, PP_ALIGN.LEFT)

    channel_x, channel_y, channel_w, channel_h = 1.2, 2.3, 14.8, 12.1
    spatial_x, spatial_y, spatial_w, spatial_h = 16.45, 2.3, 15.1, 12.1

    add_group_frame(slide, channel_x, channel_y, channel_w, channel_h, "#FFF5E6", "#DDA15E")
    add_group_frame(slide, spatial_x, spatial_y, spatial_w, spatial_h, "#F0FFF0", "#8CBF88")

    add_textbox(slide, 5.0, 2.55, 7.0, 0.7, "Channel Attention", 14, "#8A4F08", True, PP_ALIGN.CENTER)
    add_textbox(slide, 21.0, 2.55, 7.0, 0.7, "Spatial Attention", 14, "#2D6B2D", True, PP_ALIGN.CENTER)

    small_h = 1.5
    small_w = 2.5
    input_w = 2.8

    input_c = add_round_rect(slide, 1.9, 7.45, input_w, small_h, "#FFFFFF", "#666666", "Input", "#222222", 10, 1.0)
    avg_c = add_round_rect(slide, 5.3, 5.2, small_w, small_h, "#FFFFFF", "#E2A24F", "AvgPool", "#444444", 10, 1.0)
    max_c = add_round_rect(slide, 5.3, 9.35, small_w, small_h, "#FFFFFF", "#E2A24F", "MaxPool", "#444444", 10, 1.0)
    mlp_c = add_round_rect(slide, 9.0, 7.45, small_w, small_h, "#FCE5C7", "#E2A24F", "Shared MLP", "#8A4F08", 10, 1.0)
    add_c = add_operator_circle(slide, 12.35, 8.2, "+", "#FFFFFF", "#E2A24F", "#8A4F08")
    sigmoid_c = add_round_rect(slide, 13.25, 7.45, small_w, small_h, "#FCE5C7", "#E2A24F", "Sigmoid", "#8A4F08", 10, 1.0)
    mul_c = add_operator_circle(slide, 16.45, 8.2, "x", "#FFFFFF", "#E2A24F", "#8A4F08")
    out_c = add_round_rect(slide, 17.3, 7.45, input_w, small_h, "#FFFFFF", "#666666", "Refined", "#222222", 10, 1.0)

    add_polyline_arrow(slide, [(4.7, 8.2), (4.95, 8.2), (4.95, 5.95), (5.3, 5.95)], "#000000", 1.0)
    add_polyline_arrow(slide, [(4.7, 8.2), (4.95, 8.2), (4.95, 10.1), (5.3, 10.1)], "#000000", 1.0)
    add_polyline_arrow(slide, [(7.8, 5.95), (8.35, 5.95), (8.35, 8.0), (9.0, 8.0)], "#000000", 1.0)
    add_polyline_arrow(slide, [(7.8, 10.1), (8.35, 10.1), (8.35, 8.4), (9.0, 8.4)], "#000000", 1.0)
    add_straight_arrow(slide, 11.5, 8.2, 12.0, 8.2, "#000000", 1.0)
    add_straight_arrow(slide, 12.7, 8.2, 13.25, 8.2, "#000000", 1.0)
    add_straight_arrow(slide, 15.75, 8.2, 16.1, 8.2, "#000000", 1.0)
    add_straight_arrow(slide, 16.8, 8.2, 17.3, 8.2, "#000000", 1.0)

    input_s = add_round_rect(slide, 17.2, 7.45, input_w, small_h, "#FFFFFF", "#666666", "Input", "#222222", 10, 1.0)
    avg_s = add_round_rect(slide, 20.5, 5.2, small_w, small_h, "#FFFFFF", "#81B97B", "AvgPool", "#355F35", 10, 1.0)
    max_s = add_round_rect(slide, 20.5, 9.35, small_w, small_h, "#FFFFFF", "#81B97B", "MaxPool", "#355F35", 10, 1.0)
    concat_s = add_round_rect(slide, 24.15, 7.45, small_w, small_h, "#DDF3DD", "#81B97B", "Concat", "#355F35", 10, 1.0)
    conv_s = add_round_rect(slide, 27.35, 7.45, small_w, small_h, "#DDF3DD", "#81B97B", "7x7 Conv", "#355F35", 10, 1.0)
    sigmoid_s = add_round_rect(slide, 30.55, 7.45, small_w, small_h, "#DDF3DD", "#81B97B", "Sigmoid", "#355F35", 10, 1.0)
    mul_s = add_operator_circle(slide, 33.65, 8.2, "x", "#FFFFFF", "#81B97B", "#355F35")
    out_s = add_round_rect(slide, 34.45, 7.45, 2.6, small_h, "#FFFFFF", "#666666", "Output", "#222222", 10, 1.0)

    add_polyline_arrow(slide, [(20.0, 8.2), (20.2, 8.2), (20.2, 5.95), (20.5, 5.95)], "#000000", 1.0)
    add_polyline_arrow(slide, [(20.0, 8.2), (20.2, 8.2), (20.2, 10.1), (20.5, 10.1)], "#000000", 1.0)
    add_polyline_arrow(slide, [(23.0, 5.95), (23.45, 5.95), (23.45, 8.0), (24.15, 8.0)], "#000000", 1.0)
    add_polyline_arrow(slide, [(23.0, 10.1), (23.45, 10.1), (23.45, 8.4), (24.15, 8.4)], "#000000", 1.0)
    add_straight_arrow(slide, 26.65, 8.2, 27.35, 8.2, "#000000", 1.0)
    add_straight_arrow(slide, 29.85, 8.2, 30.55, 8.2, "#000000", 1.0)
    add_straight_arrow(slide, 33.05, 8.2, 33.3, 8.2, "#000000", 1.0)
    add_straight_arrow(slide, 34.0, 8.2, 34.45, 8.2, "#000000", 1.0)

    add_textbox(
        slide,
        4.6,
        14.9,
        24.5,
        0.55,
        "Parallel pooling branches are merged before attention calibration, following the standard CBAM design.",
        10,
        "#555555",
        False,
        PP_ALIGN.CENTER,
    )


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Cm(SLIDE_W_CM)
    prs.slide_height = Cm(SLIDE_H_CM)

    build_architecture_slide(prs)
    build_cbam_slide_final(prs)

    prs.save(OUTPUT_PPTX)
    print(f"Saved {OUTPUT_PPTX}")


if __name__ == "__main__":
    main()
